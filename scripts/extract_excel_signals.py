#!/usr/bin/env python3
"""Extract trading signals from the Limit Order Prototype Excel file.

Reads the Excel workbook, identifies ticker sheets (TSLA, NVDA, SPY),
extracts date and signal columns, and writes non-zero signals to CSV.

The Excel file uses two date formats:
  - datetime objects (e.g., NVDA sheet)
  - Excel serial numbers (e.g., TSLA, SPY sheets)

Both are normalized to YYYY-MM-DD strings in the output.

Usage:
    python scripts/extract_excel_signals.py [--input FILE] [--output FILE]

Output CSV format:
    Date,Ticker,Signal
    2020-02-27,NVDA,1
    2020-03-09,NVDA,1
    ...
"""

import argparse
import csv
import os
import sys
from datetime import datetime, timedelta

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)


# Excel epoch for serial number conversion (Excel uses 1900-01-01 as day 1,
# but has a leap year bug treating 1900 as a leap year, so the effective
# epoch is 1899-12-30).
EXCEL_EPOCH = datetime(1899, 12, 30)

# Ticker sheet names to process
TICKER_SHEETS = ["TSLA", "NVDA", "SPY"]

# Column indices (1-based) in each ticker sheet
DATE_COL = 1   # Column A
SIGNAL_COL = 7  # Column G


def excel_serial_to_date(serial):
    """Convert an Excel serial number to a Python date.

    Args:
        serial: Integer serial number (days since Excel epoch).

    Returns:
        datetime.date object.
    """
    return (EXCEL_EPOCH + timedelta(days=int(serial))).date()


def parse_date_value(value):
    """Normalize a date value from Excel to a date object.

    Handles both datetime objects and Excel serial numbers.

    Args:
        value: A datetime, int, or float representing a date.

    Returns:
        datetime.date object, or None if the value cannot be parsed.
    """
    if isinstance(value, datetime):
        return value.date()
    elif isinstance(value, (int, float)) and value > 0:
        return excel_serial_to_date(value)
    else:
        return None


def extract_signals(input_path, ticker_sheets=None):
    """Extract non-zero signals from the Excel workbook.

    Args:
        input_path: Path to the Excel file.
        ticker_sheets: List of sheet names to process. Defaults to
            TICKER_SHEETS.

    Returns:
        A list of dicts with keys: Date (str), Ticker (str), Signal (int).
        Also returns a summary dict mapping ticker -> signal count.
    """
    if ticker_sheets is None:
        ticker_sheets = TICKER_SHEETS

    wb = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    available_sheets = wb.sheetnames

    signals = []
    summary = {}

    for ticker in ticker_sheets:
        if ticker not in available_sheets:
            print(f"Warning: Sheet '{ticker}' not found in workbook. "
                  f"Available: {available_sheets}")
            summary[ticker] = 0
            continue

        ws = wb[ticker]
        count = 0

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=2, min_col=1, max_col=SIGNAL_COL,
                         values_only=True),
            start=2
        ):
            date_val = row[DATE_COL - 1]   # Column A (index 0)
            signal_val = row[SIGNAL_COL - 1]  # Column G (index 6)

            # Skip rows with missing data or zero signal
            if signal_val is None or signal_val == 0:
                continue

            parsed_date = parse_date_value(date_val)
            if parsed_date is None:
                print(f"Warning: Could not parse date in {ticker} "
                      f"row {row_idx}: {date_val}")
                continue

            signals.append({
                "Date": parsed_date.strftime("%Y-%m-%d"),
                "Ticker": ticker,
                "Signal": int(signal_val),
            })
            count += 1

        summary[ticker] = count

    wb.close()
    return signals, summary


def write_csv(signals, output_path):
    """Write extracted signals to a CSV file.

    Args:
        signals: List of signal dicts with Date, Ticker, Signal keys.
        output_path: Path for the output CSV file.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    with open(output_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["Date", "Ticker", "Signal"])
        writer.writeheader()
        writer.writerows(signals)


def print_summary(signals, summary, output_path):
    """Print extraction summary to stdout.

    Args:
        signals: List of signal dicts.
        summary: Dict mapping ticker -> signal count.
        output_path: Path where CSV was written.
    """
    print("\n" + "=" * 60)
    print("SIGNAL EXTRACTION SUMMARY")
    print("=" * 60)

    print(f"\nTotal signals extracted: {len(signals)}")
    print("\nSignals per ticker:")
    for ticker, count in sorted(summary.items()):
        if count > 0:
            buy_count = sum(
                1 for s in signals
                if s["Ticker"] == ticker and s["Signal"] == 1
            )
            sell_count = sum(
                1 for s in signals
                if s["Ticker"] == ticker and s["Signal"] == -1
            )
            print(f"  {ticker}: {count} signals "
                  f"({buy_count} buy, {sell_count} sell)")
        else:
            print(f"  {ticker}: {count} signals (all zeros)")

    if signals:
        print(f"\nFirst 10 rows:")
        print(f"  {'Date':<12} {'Ticker':<8} {'Signal':>6}")
        print(f"  {'-'*12} {'-'*8} {'-'*6}")
        for s in signals[:10]:
            print(f"  {s['Date']:<12} {s['Ticker']:<8} {s['Signal']:>6}")

        if len(signals) > 20:
            print(f"\n  ... ({len(signals) - 20} rows omitted) ...\n")

        print(f"Last 10 rows:")
        print(f"  {'Date':<12} {'Ticker':<8} {'Signal':>6}")
        print(f"  {'-'*12} {'-'*8} {'-'*6}")
        for s in signals[-10:]:
            print(f"  {s['Date']:<12} {s['Ticker']:<8} {s['Signal']:>6}")

    print(f"\nOutput written to: {output_path}")
    print("=" * 60)


def main():
    """Main entry point for signal extraction."""
    parser = argparse.ArgumentParser(
        description="Extract trading signals from Limit Order Prototype Excel"
    )
    parser.add_argument(
        "--input", "-i",
        default="Limit Order Prototype V9.xlsx",
        help="Path to the Excel file (default: Limit Order Prototype V9.xlsx)"
    )
    parser.add_argument(
        "--output", "-o",
        default="input/signals/limit_order_signals.csv",
        help="Output CSV path (default: input/signals/"
             "limit_order_signals.csv)"
    )
    parser.add_argument(
        "--tickers", "-t",
        nargs="+",
        default=TICKER_SHEETS,
        help=f"Ticker sheets to extract (default: {TICKER_SHEETS})"
    )
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"Error: Input file not found: {args.input}")
        sys.exit(1)

    print(f"Reading signals from: {args.input}")
    print(f"Processing tickers: {args.tickers}")

    signals, summary = extract_signals(args.input, args.tickers)

    # Sort by date then ticker for consistent output
    signals.sort(key=lambda s: (s["Date"], s["Ticker"]))

    write_csv(signals, args.output)
    print_summary(signals, summary, args.output)


if __name__ == "__main__":
    main()
